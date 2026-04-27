"""Document ingestion - loads and extracts text from files."""

import logging
import time
from datetime import datetime
from pathlib import Path
from typing import List

from app.config import SUPPORTED_EXTENSIONS
from app.models import Document

logger = logging.getLogger(__name__)


def extract_text(path: Path) -> str:
    """Extract text from a file based on its extension."""
    ext = path.suffix.lower()
    t0 = time.perf_counter()

    if ext == ".pdf":
        try:
            logger.info("Extracting PDF: %s", path.name)
            import fitz
            doc = fitz.open(str(path))
            text = "\n".join(page.get_text() for page in doc)
            elapsed = time.perf_counter() - t0
            logger.info("PDF extraction completed: %s, chars=%d, elapsed=%.3fs", path.name, len(text), elapsed)
            return text
        except Exception as e:
            elapsed = time.perf_counter() - t0
            logger.error("PDF extraction failed [%s] after %.3fs: %s", path.name, elapsed, e)
            return ""

    elif ext == ".docx":
        try:
            logger.info("Extracting DOCX: %s", path.name)
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            elapsed = time.perf_counter() - t0
            logger.info("DOCX extraction completed: %s, chars=%d, elapsed=%.3fs", path.name, len(text), elapsed)
            return text
        except Exception as e:
            elapsed = time.perf_counter() - t0
            logger.error("DOCX extraction failed [%s] after %.3fs: %s", path.name, elapsed, e)
            return ""

    elif ext == ".pptx":
        try:
            logger.info("Extracting PPTX: %s", path.name)
            from pptx import Presentation
            prs = Presentation(str(path))
            text_parts = []
            for slide_num, slide in enumerate(prs.slides, 1):
                slide_text = [f"\n[Slide {slide_num}]"]

                # Extract title (usually first text frame with title placeholder)
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        if hasattr(shape, "placeholder_format") and shape.placeholder_format.type == 1:  # Title placeholder
                            slide_text.append(f"Title: {shape.text.strip()}")
                            break

                # Extract bullet points and body text
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        # Skip if it was already captured as title
                        if slide_text and slide_text[-1].startswith("Title:") and shape.text.strip() == slide_text[-1][7:].strip():
                            continue

                        # Check if it's a text frame with paragraphs (bullets)
                        if hasattr(shape, "text_frame"):
                            for para in shape.text_frame.paragraphs:
                                if para.text.strip():
                                    bullet_marker = "• " if para.level > 0 else ""
                                    slide_text.append(f"{bullet_marker}{para.text.strip()}")
                        else:
                            # Regular text shape
                            slide_text.append(shape.text.strip())

                text_parts.append("\n".join(slide_text))

            text = "\n".join(text_parts)
            elapsed = time.perf_counter() - t0
            logger.info("PPTX extraction completed: %s, chars=%d, elapsed=%.3fs", path.name, len(text), elapsed)
            return text
        except Exception as e:
            elapsed = time.perf_counter() - t0
            logger.error("PPTX extraction failed [%s] after %.3fs: %s", path.name, elapsed, e)
            return ""

    elif ext == ".html":
        try:
            logger.info("Extracting HTML: %s", path.name)
            from html.parser import HTMLParser

            class _TextExtractor(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self._parts: list[str] = []
                    self._skip = False

                def handle_starttag(self, tag, attrs):
                    if tag in ("script", "style"):
                        self._skip = True

                def handle_endtag(self, tag):
                    if tag in ("script", "style"):
                        self._skip = False
                    if tag in ("p", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "br"):
                        self._parts.append("\n")

                def handle_data(self, data):
                    if not self._skip:
                        self._parts.append(data)

            raw = path.read_text(encoding="utf-8", errors="ignore")
            extractor = _TextExtractor()
            extractor.feed(raw)
            text = "".join(extractor._parts).strip()
            elapsed = time.perf_counter() - t0
            logger.info("HTML extraction completed: %s, chars=%d, elapsed=%.3fs", path.name, len(text), elapsed)
            return text
        except Exception as e:
            elapsed = time.perf_counter() - t0
            logger.error("HTML extraction failed [%s] after %.3fs: %s", path.name, elapsed, e)
            return ""

    else:  # .txt, .md, etc.
        try:
            logger.info("Extracting text file: %s", path.name)
            text = path.read_text(encoding="utf-8", errors="ignore")
            elapsed = time.perf_counter() - t0
            logger.info("Text extraction completed: %s, chars=%d, elapsed=%.3fs", path.name, len(text), elapsed)
            return text
        except Exception as e:
            elapsed = time.perf_counter() - t0
            logger.error("Text extraction failed [%s] after %.3fs: %s", path.name, elapsed, e)
            return ""


def load_documents(folder: Path) -> List[Document]:
    """Load all supported documents from a folder."""
    if not folder.exists():
        raise FileNotFoundError(f"Data folder not found: {folder}")

    documents: List[Document] = []

    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        content = extract_text(path)
        if not content.strip():
            logger.warning(f"Skipped (empty): {path.name}")
            continue

        modified_time = datetime.fromtimestamp(path.stat().st_mtime)
        doc = Document(
            source=path.name,
            path=str(path),
            text=content,
            doc_type=path.suffix.lower().lstrip("."),
            modified_time=modified_time,
        )
        documents.append(doc)
        logger.info(f"Loaded: {path.name} ({len(content)} chars)")

    if not documents:
        raise RuntimeError(
            f"No indexable content found in {folder}.\n"
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    logger.info(f"Total documents loaded: {len(documents)}")
    return documents
